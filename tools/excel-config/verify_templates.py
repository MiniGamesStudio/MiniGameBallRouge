# -*- coding: utf-8 -*-
"""
校验 tools/excel-config/*.xlsx 是否能被 tools/excel-exporter 正确导出，
以及表与表之间的一致性（引用、枚举范围、行模板、权重数组长度）。

本脚本复刻 ExcelReader.ts 的规则：
  - 第 1 行字段名、第 2 行类型、第 3 行注释、第 4 行起数据
  - 跳过空行、跳过第一列以 # 开头的注释行
  - int/float 必须是数字；bool 可为真假值；array:X 必须是 JSON 数组字符串
  - 类型标注必须是 int/float/bool/string/enum:X/array:X

用法： python verify_templates.py
退出码：0 = 全部通过，1 = 有错误
"""

import json
import os
import re
import sys
from openpyxl import load_workbook

DIR = os.path.dirname(os.path.abspath(__file__))
VALID_TYPES = ("int", "float", "bool", "string")
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def to_pascal_case(name):
    """复刻 SchemaGenerator/ExportPipeline 的 toPascalCase：按 _ - 空格 分段，每段首字母大写、其余小写。"""
    parts = [p for p in re.split(r"[_\-\s]+", name) if p]
    return "".join(p[0].upper() + p[1:].lower() for p in parts)


def underscore_form(name):
    """给出 camelCase → Underscore_Form 的建议写法。"""
    return re.sub(r"(?<!^)(?=[A-Z])", "_", name)

errors = []
warnings = []
info = []


def err(msg):
    errors.append(msg)


def warn(msg):
    warnings.append(msg)


def is_valid_type(t):
    if t in VALID_TYPES:
        return True
    if t.startswith("enum:") and len(t) > 5:
        return True
    if t.startswith("array:"):
        inner = t[6:]
        return inner in VALID_TYPES or (inner.startswith("enum:") and len(inner) > 5)
    return False


def check_cell(value, t, where):
    """复刻 ExcelReader.validateCellType（JS 语义 → Python）"""
    if value is None or value == "":
        return True
    if t == "int":
        if isinstance(value, bool):
            return False
        if isinstance(value, int):
            return True
        if isinstance(value, float):
            return float(value).is_integer()
        return False
    if t == "float":
        if isinstance(value, bool):
            return False
        return isinstance(value, (int, float))
    if t == "bool":
        if isinstance(value, bool):
            return True
        if isinstance(value, (int, float)):
            return value in (0, 1)
        if isinstance(value, str):
            return value.lower() in ("true", "false", "0", "1")
        return False
    if t == "string":
        return True
    if t.startswith("enum:"):
        return isinstance(value, (int, float, str))
    if t.startswith("array:"):
        if isinstance(value, list):
            return True
        if isinstance(value, str):
            try:
                return isinstance(json.loads(value), list)
            except Exception:
                return False
        return False
    return False


def read_sheet(ws, filename):
    """返回 (fields, rows)；fields = [(name, type, comment)]"""
    rows_all = list(ws.iter_rows(values_only=True))
    if len(rows_all) < 3:
        err(f"{filename}/{ws.title}: 少于 3 行表头")
        return None, None
    names, types, comments = rows_all[0], rows_all[1], rows_all[2]

    fields = []
    for i, (n, t, c) in enumerate(zip(names, types, comments)):
        if n is None or str(n).strip() == "":
            continue
        n = str(n).strip()
        t = "" if t is None else str(t).strip()
        if not IDENT.match(n):
            err(f"{filename}/{ws.title}: 字段名 '{n}' 不是合法的标识符（FlatBuffers 需要 ASCII）")
        if not is_valid_type(t):
            err(f"{filename}/{ws.title}: 字段 '{n}' 的类型标注 '{t}' 不被导出器支持")
        fields.append((n, t, "" if c is None else str(c)))
    if not fields:
        err(f"{filename}/{ws.title}: 没有解析到任何字段")

    data = []
    for idx, row in enumerate(rows_all[3:], start=4):
        if row is None or all(v is None or v == "" for v in row):
            continue                      # 空行跳过
        first = row[0]
        if isinstance(first, str) and first.startswith("#"):
            continue                      # 注释行跳过
        row = list(row[:len(fields)])
        while len(row) < len(fields):
            row.append(None)
        for i, (n, t, _) in enumerate(fields):
            if not check_cell(row[i], t, n):
                err(f"{filename}/{ws.title}: 第 {idx} 行 '{n}' 的值 {row[i]!r} 与类型 '{t}' 不匹配（会导致整表导出中止）")
        data.append(row)
    return fields, data


def col(fields, name):
    for i, (n, _, _) in enumerate(fields):
        if n == name:
            return i
    return None


def as_list(v):
    if v is None or v == "":
        return None
    if isinstance(v, list):
        return v
    try:
        out = json.loads(v)
        return out if isinstance(out, list) else None
    except Exception:
        return None


def main():
    tables = {}   # (file, sheet) -> (fields, rows)
    table_names = {}   # 导出后的表名 -> 来源，用于查重

    files = ["player.xlsx", "enemy.xlsx", "level.xlsx", "skill.xlsx",
             "drop.xlsx", "meta.xlsx", "skin.xlsx", "tuning.xlsx"]
    for fn in files:
        path = os.path.join(DIR, fn)
        if not os.path.exists(path):
            err(f"缺少配置文件：{fn}")
            continue
        wb = load_workbook(path, data_only=True)
        for sheet in wb.sheetnames:
            if not IDENT.match(sheet):
                err(f"{fn}: 工作表名 '{sheet}' 不是合法标识符（会生成非法表名）")
                continue
            # 工作表名规范：每段都要是「首字母大写 + 其余小写」，多词用 _ 分隔。
            # 反例 camelCase 的 "BossPhase"：导出器会压平成 "Bossphase"。
            tname = to_pascal_case(sheet)
            parts = [p for p in re.split(r"[_\-\s]+", sheet) if p]
            bad = [p for p in parts if not re.fullmatch(r"[A-Za-z][a-z0-9]*", p)]
            if bad:
                err(f"{fn}: 工作表名 '{sheet}' 里 {bad} 是 camelCase 片段，导出后表名会被压平成 "
                    f"'{tname}'；请改成 '{underscore_form(sheet)}'")
            if tname in table_names:
                err(f"表名冲突：'{tname}' 同时来自 {table_names[tname]} 与 {fn}/{sheet}")
            table_names[tname] = f"{fn}/{sheet}"
            fields, rows = read_sheet(wb[sheet], fn)
            if fields is not None:
                tables[(fn, sheet)] = (fields, rows)
                info.append(f"{fn:<14} {sheet:<12} 字段 {len(fields):>2} 个，数据 {len(rows):>2} 行  → 表名 {tname}")
        wb.close()

    # ── 主键唯一性 ────────────────────────────────────────────
    for (fn, sheet), (fields, rows) in tables.items():
        i = col(fields, "id")
        if i is None:
            continue
        seen = set()
        for r in rows:
            key = r[i]
            if key in seen:
                err(f"{fn}/{sheet}: 主键 '{key}' 重复")
            seen.add(key)

    # ── 枚举范围 ─────────────────────────────────────────────
    enum_ranges = {
        "quality": (0, 5), "type": (0, 3), "shape": (0, 5), "attackKind": (0, 4),
        "kind": (0, 3), "rarity": (0, 3), "sourceKind": (0, 3), "unlock": (0, 2),
        "mode": (0, 1), "op": (0, 1), "effectOp": (0, 1),
    }
    for (fn, sheet), (fields, rows) in tables.items():
        for name, (lo, hi) in enum_ranges.items():
            i = col(fields, name)
            if i is None:
                continue
            for r in rows:
                v = r[i]
                if isinstance(v, (int, float)) and not (lo <= v <= hi):
                    err(f"{fn}/{sheet}: {name} = {v} 超出范围 [{lo},{hi}]")

    # ── 概率范围 ─────────────────────────────────────────────
    for (fn, sheet), (fields, rows) in tables.items():
        i = col(fields, "chance")
        if i is None:
            continue
        for r in rows:
            v = r[i]
            if isinstance(v, (int, float)) and not (0.0 <= v <= 1.0):
                err(f"{fn}/{sheet}: chance = {v} 不在 0~1")

    # ── 敌人：类型 × 体型 映射矩阵（§7.5）────────────────────
    ef, er = tables[("enemy.xlsx", "Enemy")]
    iq, it, ish, i_id = col(ef, "quality"), col(ef, "type"), col(ef, "shape"), col(ef, "id")
    allowed = {
        0: {0, 1}, 1: {0, 1},   # 单格 / 双格横：普通、精英
        2: {0, 1},              # 双格竖
        3: {2},                 # 四格：小BOSS
        4: {2, 3},              # 六格：小BOSS / 大BOSS
        5: {3},                 # 八格：大BOSS
    }
    enemy_ids = set()
    for r in er:
        eid, shape, etype = r[i_id], r[ish], r[it]
        enemy_ids.add(eid)
        if shape in allowed and etype not in allowed[shape]:
            err(f"enemy.xlsx/Enemy: {eid} 的体型 {shape} 与类型 {etype} 不在允许矩阵内（§7.5）")
        info.append(f"  敌人 {eid:<20} 品质{r[iq]} 类型{etype} 体型{shape}")

    # ── 关卡：行模板与引用（含多格占位的覆盖校验）───────────
    # 体型 → (colSpan, rowSpan)，与 §7.4 一致
    SHAPE_SPAN = {0: (1, 1), 1: (2, 1), 2: (1, 2), 3: (2, 2), 4: (3, 2), 5: (4, 2)}
    shape_of = {r[i_id]: r[ish] for r in er}

    wf, wr = tables[("level.xlsx", "Wave")]
    rf, rr = tables[("level.xlsx", "Row")]
    i_lvl_w, i_idx_w = col(wf, "levelId"), col(wf, "waveIndex")
    i_lvl_r, i_idx_r = col(rf, "levelId"), col(rf, "waveIndex")
    i_rowidx_r, i_tpl = col(rf, "rowIndex"), col(rf, "template")

    waves = {(r[i_lvl_w], r[i_idx_w]): r for r in wr}
    for r in wr:
        q = as_list(r[col(wf, "qualityWeights")])
        t = as_list(r[col(wf, "typeWeights")])
        s = as_list(r[col(wf, "shapeWeights")])
        if q is not None and len(q) != 6:
            err(f"level.xlsx/Wave: {r[0]} 的 qualityWeights 必须是 6 项，实为 {len(q)}")
        if t is not None and len(t) != 4:
            err(f"level.xlsx/Wave: {r[0]} 的 typeWeights 必须是 4 项，实为 {len(t)}")
        if s is not None and len(s) != 6:
            err(f"level.xlsx/Wave: {r[0]} 的 shapeWeights 必须是 6 项，实为 {len(s)}")
        if r[col(wf, "mode")] == 1 and (q is None or sum(q) == 0):
            err(f"level.xlsx/Wave: {r[0]} 是随机波但品质权重全为 0")

    # 按 (关卡, 波次) 分组做网格覆盖校验：
    #   defId = 锚点格；'-' = 被锚点覆盖的格（同行左侧的横版，或上一行的竖版/大怪）；'_' = 空列
    groups = {}
    for r in rr:
        groups.setdefault((r[i_lvl_r], r[i_idx_r]), []).append(r)

    for key, rows in groups.items():
        if key not in waves:
            err(f"level.xlsx/Row: 波次 {key} 没有对应的 Wave 记录")
        label = f"{key[0]}/W{key[1]}"
        grid = {}          # (rowIdx, col) -> 'anchor' | 'occupied'
        claimed = set()    # 被 '-' 明确声明的占位格
        for r in sorted(rows, key=lambda x: x[i_rowidx_r]):
            ridx = r[i_rowidx_r]
            slots = str(r[i_tpl]).split("|")
            if len(slots) != 5:
                err(f"level.xlsx/Row: {r[0]} 的模板有 {len(slots)} 个槽位，应为 5（colCount）")
                continue
            for k, slot in enumerate(slots):
                if slot == "-":
                    if grid.get((ridx, k)) != "occupied":
                        err(f"level.xlsx/Row: {r[0]} 第 {k} 槽位是 '-'，但没有任何锚点覆盖它")
                    else:
                        claimed.add((ridx, k))
                    continue
                if slot in ("_", ""):
                    if (ridx, k) in grid:
                        err(f"level.xlsx/Row: {r[0]} 第 {k} 槽位写的是空列，但它已被锚点占用")
                    continue
                # 锚点格
                if slot not in enemy_ids:
                    err(f"level.xlsx/Row: {r[0]} 引用了不存在的敌人 '{slot}'")
                    continue
                if (ridx, k) in grid:
                    err(f"level.xlsx/Row: {r[0]} 第 {k} 槽位的锚点 '{slot}' 与已有锚点重叠")
                cs, rs = SHAPE_SPAN.get(shape_of.get(slot, 0), (1, 1))
                if k + cs > 5:
                    err(f"level.xlsx/Row: {r[0]} 的 '{slot}' 从第 {k} 列展开 {cs} 列，超出 5 列棋盘")
                grid[(ridx, k)] = "anchor"
                for dr in range(rs):
                    for dc in range(cs):
                        if dr == 0 and dc == 0:
                            continue
                        grid[(ridx + dr, k + dc)] = "occupied"
        # 覆盖完整性：锚点压住的每一格都必须由某一行的 '-' 明确写出
        missing = sorted(c for c, v in grid.items() if v == "occupied" and c not in claimed)
        for (ridx, k) in missing:
            err(f"level.xlsx/Row: {label} 第 {ridx} 行第 {k} 列被多格敌人占用，但没有任何行把它写成 '-'")
        for (ridx, k) in sorted(claimed):
            if grid.get((ridx, k)) != "occupied":
                err(f"level.xlsx/Row: {label} 第 {ridx} 行第 {k} 列写成 '-' 但没有锚点覆盖")
        anchors = sum(1 for v in grid.values() if v == "anchor")
        info.append(f"  关卡 {label:<12} 行 {len(rows)}，锚点 {anchors}，占位格 {len(claimed)}")

    # ── 技能：效果 / 进化 / 融合 引用 ────────────────────────
    sf, sr = tables[("skill.xlsx", "Skill")]
    i_sid = col(sf, "id")
    skill_ids = {r[i_sid] for r in sr}
    eff, effr = tables[("skill.xlsx", "Skill_Effect")]
    i_esid = col(eff, "skillId")
    for r in effr:
        if r[i_esid] not in skill_ids:
            err(f"skill.xlsx/SkillEffect: 引用了不存在的技能 '{r[i_esid]}'")

    for r in sr:
        sid = r[i_sid]
        ev = r[col(sf, "evolveInto")]
        if ev not in (None, "", "无") and ev not in skill_ids:
            err(f"skill.xlsx/Skill: {sid} 的 evolveInto 指向不存在的技能 '{ev}'")
        fus = as_list(r[col(sf, "fusePartners")])
        for entry in (fus or []):
            parts = str(entry).split(":")
            if len(parts) != 3:
                err(f"skill.xlsx/Skill: {sid} 的 fusePartners 条目 '{entry}' 格式应为 对方ID:产物ID:消耗")
                continue
            other, into, cost = parts
            if other not in skill_ids:
                err(f"skill.xlsx/Skill: {sid} 的融合对象 '{other}' 不存在")
            if into not in skill_ids:
                err(f"skill.xlsx/Skill: {sid} 的融合产物 '{into}' 不存在")
            if int(cost) < 0:
                err(f"skill.xlsx/Skill: {sid} 的融合消耗不能为负")
        # 融合必须双向登记
        for entry in (fus or []):
            other, into, _ = str(entry).split(":")
            other_row = next((x for x in sr if x[i_sid] == other), None)
            back = as_list(other_row[col(sf, "fusePartners")]) if other_row else None
            if not any(str(e).split(":")[1] == into and str(e).split(":")[0] == sid for e in (back or [])):
                warn(f"skill.xlsx/Skill: {sid} ↔ {other} 的融合未双向登记（应互为伙伴）")

    # 每个技能至少有一条效果
    has_eff = {r[i_esid] for r in effr}
    for r in sr:
        if r[i_sid] not in has_eff:
            warn(f"skill.xlsx/Skill: 技能 {r[i_sid]} 没有任何效果行")

    # ── 掉落表引用 ───────────────────────────────────────────
    df, dr = tables[("drop.xlsx", "Drop")]
    for r in dr:
        if r[col(df, "sourceKind")] == 3:
            sid = r[col(df, "sourceId")]
            if sid not in enemy_ids:
                err(f"drop.xlsx/Drop: {r[0]} 指向不存在的敌人 '{sid}'")

    # ── 输出 ─────────────────────────────────────────────────
    print("=" * 78)
    print("配置表校验报告")
    print("=" * 78)
    for line in info:
        if line.startswith("  "):
            print(line)
    print("-" * 78)
    for line in info:
        if not line.startswith("  "):
            print(line)
    print("-" * 78)
    print(f"警告 {len(warnings)} 条：")
    for w in warnings:
        print("  [WARN] " + w)
    print(f"错误 {len(errors)} 条：")
    for e in errors:
        print("  [ERROR] " + e)
    print("=" * 78)
    if errors:
        print("结果：不通过（存在错误，导出会失败或被引用校验拦截）")
        return 1
    print("结果：通过（表头/类型/引用/枚举/权重/行模板 全部合法）")
    return 0


if __name__ == "__main__":
    sys.exit(main())