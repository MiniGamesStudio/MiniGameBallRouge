# -*- coding: utf-8 -*-
"""把新素材写进配置表（128 格 + 品质底图 + 怪物/BOSS/玩家贴图）。
⚠️ 这是对 .xlsx 的直接补丁：build_templates.py 才是生成源，等同步盘问题解决后要把同样的改动补回脚本。
"""
import openpyxl, os, sys

d = os.path.dirname(os.path.abspath(__file__))
log = []

def rows_of(ws, header_row, first_data_row):
    return first_data_row

def find_header(ws, field, max_scan=5):
    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value == field:
                return r, c
    return None, None

# ── 1. enemy.xlsx：Quality 加 baseArt 列；Enemy.asset 换真实贴图名 ──
p = os.path.join(d, 'enemy.xlsx')
wb = openpyxl.load_workbook(p)
ws = wb['Quality']
hr, hc = find_header(ws, 'colorHex')
if hr is None:
    sys.exit('ERROR: Quality 找不到 colorHex 列')
if find_header(ws, 'baseArt')[0] is None:
    ws.insert_cols(hc + 1)
    ws.cell(row=hr, column=hc + 1, value='baseArt')
    ws.cell(row=hr + 1, column=hc + 1, value='string')
    ws.cell(row=hr + 2, column=hc + 1, value='品质底图贴图名（一格一张，128×128）')
    arts = ['white', 'green', 'blue', 'purple', 'yellow', 'red']
    r0 = hr + 3
    for i, a in enumerate(arts):
        ws.cell(row=r0 + i, column=hc + 1, value=a)
    log.append('Quality +baseArt 列(第%d列) 6 行' % (hc + 1))
ws.cell(row=hr + 2, column=hc, value='染色色值（UI / 掉落物染色仍可用）')

ws2 = wb['Enemy']
hr2, hc2 = find_header(ws2, 'asset')
ws2.cell(row=hr2 + 2, column=hc2, value='怪物贴图名（monster_0001-0010 普通/精英；Boss_001-003 四/六/八格）')
new_asset = {
    'e_slime_white': 'monster_0001', 'e_slime_green': 'monster_0002', 'e_archer_green': 'monster_0003',
    'e_gunner_blue': 'monster_0004', 'e_laser_purple': 'monster_0005', 'e_charger_gold': 'monster_0006',
    'e_miniboss_blue4': 'Boss_001', 'e_boss_red6': 'Boss_002', 'e_boss_red8': 'Boss_003',
}
n = 0
for r in range(hr2 + 3, ws2.max_row + 1):
    eid = ws2.cell(row=r, column=1).value
    if eid in new_asset:
        ws2.cell(row=r, column=hc2, value=new_asset[eid]); n += 1
log.append('Enemy.asset 更新 %d 行' % n)
wb.save(p)

# ── 2. player.xlsx：贴图名 + 128 格半径/速度 ──
p = os.path.join(d, 'player.xlsx')
wb = openpyxl.load_workbook(p)
ws = wb['Player']
hr, _ = find_header(ws, 'id')
vals = {'sprite': 'player_001', 'visualRadius': 56, 'hitRadius': 45, 'moveSpeed': 512,
        'bulletSpeed': 1440, 'bulletRadius': 16, 'catchRadius': 61, 'magnetRadius': 192}
cols = {}
for name in vals:
    r, c = find_header(ws, name)
    cols[name] = c
    ws.cell(row=r + 2, column=c, value=('玩家贴图名（player_001）' if name == 'sprite'
             else ('受击判定半径 px（建议 45，越小越耐玩）' if name == 'hitRadius' else ws.cell(row=r + 2, column=c).value)))
for r in range(hr + 3, ws.max_row + 1):
    if ws.cell(row=r, column=1).value:
        for name, v in vals.items():
            ws.cell(row=r, column=cols[name], value=v)
log.append('Player 更新: ' + ', '.join('%s=%s' % (k, v) for k, v in vals.items()))
wb.save(p)

# ── 3. tuning.xlsx：格尺寸与 px 数值 ──
p = os.path.join(d, 'tuning.xlsx')
wb = openpyxl.load_workbook(p)
ws = wb['Tuning']
hr, _ = find_header(ws, 'key')
upd = {'cellSize': 128, 'bottomLineOffset': 160, 'playerSpawnOffset': 192,
       'diveSpeed': 1120, 'diveHitRadius': 72, 'dropScatterRadius': 26}
found = {}
for r in range(hr + 3, ws.max_row + 1):
    k = ws.cell(row=r, column=1).value
    if k in upd:
        old = ws.cell(row=r, column=2).value
        ws.cell(row=r, column=2, value=upd[k]); found[k] = (old, upd[k])
        if k == 'cellSize':
            ws.cell(row=r, column=3, value='场地')
            ws.cell(row=r, column=4, value='格尺寸 px（= 品质底图素材尺寸，体型与散落距离的基准）')
        if k == 'dropScatterRadius':
            ws.cell(row=r, column=4, value='散落半径 px（0.2 格 = 25.6）')
log.append('Tuning 更新: ' + ', '.join('%s %s→%s' % (k, v[0], v[1]) for k, v in found.items()))
keys = [ws.cell(row=r, column=1).value for r in range(hr + 3, ws.max_row + 1)]
if 'artFitMargin' not in keys:
    nr = ws.max_row + 1
    ws.cell(row=nr, column=1, value='artFitMargin')
    ws.cell(row=nr, column=2, value=0.92)
    ws.cell(row=nr, column=3, value='美术')
    ws.cell(row=nr, column=4, value='怪物/玩家图在占格内的留白系数（contain 适配）')
    log.append('Tuning +artFitMargin 0.92')
wb.save(p)

print('OK')
for l in log:
    print('  -', l)